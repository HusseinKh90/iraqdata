"""
pipeline.parsers.kfi_timeline
=============================
The month columns of the CBI Key Financial Indicators workbook, read the same way by both parsers
that use it (cbi_cpi and cbi_money_banking).

Row 1 holds the year at the first column of each year, row 2 the month ("Jan.", "Feb", "Sept.",
"March" ...). Monthly columns start at column 3 and run to the "The week ending" block, which
holds the latest weekly snapshot and is not part of the monthly series.

Every column up to that block must be a recognised month, and the months must follow one another
without a gap or a repeat. A header that is not recognised (a new spelling), a blank one, a
forgotten year header or a skipped month stops the run with the cell named: a column that was
skipped silently would drop its month from the published series without any other check noticing.
"""

import re
from typing import Dict, Optional

from openpyxl.utils import get_column_letter

import source_io

FIRST_MONTH = "2004-01"  # the sample starts here; the workbook's December 2003 column is not read
WEEKLY_MARKER = "week ending"
FIRST_COLUMN = 3

_NAMES = ("january", "february", "march", "april", "may", "june", "july", "august", "september",
          "october", "november", "december")
# full names, three-letter abbreviations and the four-letter "sept"; a trailing "." is ignored
_MONTHS = {**{name: i for i, name in enumerate(_NAMES, start=1)},
           **{name[:3]: i for i, name in enumerate(_NAMES, start=1)}, "sept": 9}
_YEAR = re.compile(r"[0-9]{4}")


def month_of(label) -> Optional[int]:
    """1 to 12 for a month header such as 'Jan.', 'Feb ', 'Sept.' or 'March'; None for anything else."""
    if not isinstance(label, str):
        return None
    return _MONTHS.get(label.strip().lower().rstrip(".").strip())


def _year_of(value) -> Optional[int]:
    """A year header: a whole number, or four digits as text (Arabic-Indic digits are read as digits)."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if value == int(value) and 1900 <= value <= 2100 else None
    text = str(value).strip().translate(source_io.ARABIC_INDIC)
    return int(text) if _YEAR.fullmatch(text) else None


def resolve(ws, path, sheet: Optional[str] = None) -> Dict[int, str]:
    """
    {column: 'YYYY-MM'} for the monthly columns of the worksheet from FIRST_MONTH on. Raises a
    source_io.ParseError naming the cell for an unrecognised, blank or out-of-order header.
    """
    sheet = sheet or ws.title
    columns = []
    for col in range(FIRST_COLUMN, ws.max_column + 1):
        year, month = ws.cell(1, col).value, ws.cell(2, col).value
        if isinstance(year, str) and WEEKLY_MARKER in year.lower():
            break
        columns.append((col, year, month))
    while columns and columns[-1][1] is None and columns[-1][2] is None:
        columns.pop()  # empty columns after the last month

    dates: Dict[int, str] = {}
    year: Optional[int] = None
    previous: Optional[int] = None  # months since year 0 of the last column read
    for col, year_cell, month_cell in columns:
        letter = get_column_letter(col)
        if year_cell is not None:
            year = _year_of(year_cell)
            if year is None:
                raise source_io.ParseError(path, "a four-digit year (or the 'The week ending' block)",
                                           repr(year_cell), sheet=sheet, location=f"cell {letter}1")
        month = month_of(month_cell)
        if month is None:
            raise source_io.ParseError(path, "a month name such as Jan., Feb or Sept.", repr(month_cell),
                                       sheet=sheet, location=f"cell {letter}2")
        if year is None:
            raise source_io.ParseError(path, "a year header before the first month", repr(month_cell),
                                       sheet=sheet, location=f"cell {letter}1")
        index = year * 12 + month - 1
        if previous is not None and index != previous + 1:
            expected = f"{(previous + 1) // 12}-{(previous + 1) % 12 + 1:02d}"
            raise source_io.ParseError(path, f"the month after the previous column, {expected}",
                                       f"{year}-{month:02d}", sheet=sheet, location=f"cell {letter}2")
        previous = index
        date = f"{year}-{month:02d}"
        if date >= FIRST_MONTH:
            dates[col] = date
    if not dates:
        raise source_io.ParseError(path, f"monthly timeline columns from {FIRST_MONTH}", "none", sheet=sheet)
    return dates
