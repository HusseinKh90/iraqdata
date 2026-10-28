"""
pipeline.parsers.parse_cbi_money_banking
========================================
Parser for Family 1 (Money & Banking) from the Central Bank of Iraq (CBI)
Key Financial Indicators workbook (https://cbi.iq/news/view/94).

Extracts 10 monthly money and banking indicators:
1. money.currency_outside_banks (BILLION_IQD)
2. money.m0 (BILLION_IQD)
3. money.m1 (BILLION_IQD)
4. money.m2 (BILLION_IQD)
5. money.fx_deposits (BILLION_IQD)
6. money.reserves_official (BILLION_IQD)
7. money.policy_rate (PERCENT)
8. money.bank_deposit_rate (PERCENT)
9. money.bank_lending_rate (PERCENT)
10. money.required_reserves (BILLION_IQD)

Quality & Schema Guarantees:
- Single-frequency rule: Produces purely monthly observations (YYYY-MM).
- Conforms to standard 7-column Canonical Observation Schema.
- The money supply ordering (M1 < M2, currency <= M0 and M1) is checked by the gatekeeper,
  pipeline/validate.py, not here.
- Output: pipeline/data/garden/money_banking_monthly.parquet
"""

import re
import sys
import datetime
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd

# Add src to sys.path so iraqdata package modules can be imported
repo_root = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(repo_root / "src"))
sys.path.insert(0, str(repo_root / "pipeline"))

from iraqdata.schema import CANONICAL_COLUMNS
import flags  # noqa: E402
import kfi_timeline  # noqa: E402
import source_io  # noqa: E402

# Indicator definitions: (series_id, row in the 2026-10 workbook, expected label, unit).
# Rows are found by label: the column B label, normalised by normalise_label,
# must equal the expected label exactly, in exactly one row above the "Source:" line. The row number
# only appears in error messages.
INDICATORS = [
    ("money.policy_rate", 33, "policy rate", "PERCENT"),
    ("money.bank_deposit_rate", 58, "id bank deposits one year fixed", "PERCENT"),
    ("money.bank_lending_rate", 60, "id bank loans 1 5 year", "PERCENT"),
    ("money.m0", 63, "monetary base end of period in billions of iraqi dinars", "BILLION_IQD"),
    ("money.reserves_official", 67, "official reserve", "BILLION_IQD"),
    ("money.currency_outside_banks", 102, "currency outside banks", "BILLION_IQD"),
    ("money.m1", 109, "m1", "BILLION_IQD"),
    ("money.m2", 110, "m2", "BILLION_IQD"),
    ("money.fx_deposits", 113, "foreign currency deposits", "BILLION_IQD"),
    ("money.required_reserves", 116, "required reserves", "BILLION_IQD"),
]

REALTIME_START = "2026-10-28"
SHEET = "Sheet1"  # the only sheet of the Key Financial Indicators workbook
NO_VALUE = {"-", "…"}  # cell texts that mean "no value"; blank cells are no value too
DATED_NAME_INDICATORS = re.compile(r"cbi_key_financial_indicators_\d{4}-\d{2}-\d{2}\.xlsx")


compute_file_hash = source_io.file_hash  # 16 hex characters of the file's SHA-256


def normalise_label(value) -> str:
    """
    Column B label for exact comparison: lowercase, every run of other characters one space, a
    leading one-letter list marker ("a - ") and trailing one- or two-digit footnote numbers dropped.
    "   a - Required Reserves 3" -> "required reserves"; "Of which: Deposits component of M2" stays distinct from "m2".
    """
    tokens = re.sub(r"[^0-9a-z]+", " ", str(value or "").lower()).split()
    if len(tokens) > 1 and len(tokens[0]) == 1 and tokens[0].isalpha():
        tokens = tokens[1:]
    while len(tokens) > 1 and tokens[-1].isdigit() and len(tokens[-1]) <= 2:
        tokens.pop()
    return " ".join(tokens)


def resolve_indicators_input(raw_cbi_dir: Path) -> Optional[Path]:
    """
    Picks the CBI Key Financial Indicators workbook to parse: the newest dated download
    (cbi_key_financial_indicators_<YYYY-MM-DD>.xlsx), else the legacy archived key_financial_indicators.xlsx, else None.
    """
    dated = sorted(p for p in raw_cbi_dir.glob("cbi_key_financial_indicators_*.xlsx") if DATED_NAME_INDICATORS.fullmatch(p.name))
    if dated:
        return dated[-1]
    legacy = raw_cbi_dir / "key_financial_indicators.xlsx"
    return legacy if legacy.exists() else None


def download_latest_indicators(
    destination_dir: Path,
    today: Optional[date] = None,
    manifest_path: Optional[Path] = None,
    manifest_root: Optional[Path] = None,
) -> Path:
    """
    Finds the latest Key Financial Indicators workbook on https://cbi.iq/news/view/94 and archives it as a new
    dated file, cbi_key_financial_indicators_<YYYY-MM-DD>.xlsx, in destination_dir, then adds it to the raw-file
    manifest. Never overwrites an existing file; TLS is verified and robots.txt respected.
    """
    import downloads
    import raw_audit

    target = Path(destination_dir) / f"cbi_key_financial_indicators_{(today or datetime.datetime.now(datetime.timezone.utc).date()).isoformat()}.xlsx"
    if target.exists():
        raise FileExistsError(target)

    page_url = "https://cbi.iq/news/view/94"
    response = downloads.polite_get(page_url, timeout=20, max_bytes=downloads.PAGE_MAX_BYTES)
    response.raise_for_status()

    # The workbook is chosen by its content (the sheet the parser reads), never by its position on
    # the page; an error page is never archived
    file_url, content = downloads.pick_workbook(page_url, response.text, sheets=(SHEET,))
    downloads.save_new_file(target, content)
    print(f"Downloaded indicators from {file_url} ({len(content):,} bytes) -> {target}")
    raw_audit.manifest_add(
        target,
        manifest_path if manifest_path is not None else raw_audit.MANIFEST_PATH,
        root=manifest_root if manifest_root is not None else repo_root,
    )
    return target


def parse_cbi_money_banking(
    workbook_path: Optional[Path] = None,
    output_parquet_path: Optional[Path] = None
) -> pd.DataFrame:
    """
    Parses Key Financial Indicators workbook and outputs canonical monthly
    Money & Banking dataset into pipeline/data/garden/money_banking_monthly.parquet.
    """
    if workbook_path is None:
        raw_cbi_dir = repo_root / "pipeline" / "raw" / "cbi"
        workbook_path = resolve_indicators_input(raw_cbi_dir)
        if workbook_path is None:
            workbook_path = download_latest_indicators(raw_cbi_dir)

    if output_parquet_path is None:
        output_parquet_path = repo_root / "pipeline" / "data" / "garden" / "money_banking_monthly.parquet"

    wb = source_io.open_workbook(workbook_path, [SHEET], data_only=True)
    ws = wb[SHEET]  # read by name, never the workbook's active sheet
    source_hash = compute_file_hash(workbook_path)
    report = source_io.Report("cbi_money_banking")

    # 1. Resolve timeline across columns (Row 1 = Year, Row 2 = Month); every month header must be
    #    recognised and the months consecutive (kfi_timeline)
    col_dates: Dict[int, str] = kfi_timeline.resolve(ws, workbook_path)

    print(f"Extracted {len(col_dates)} monthly columns: from {min(col_dates.values())} to {max(col_dates.values())}.")

    # 2. Resolve indicator rows by exact label within the main table (above the "Source:" line;
    #    the memorandum block below it repeats M1 and M2 as growth rates)
    end_row = next((r for r in range(1, ws.max_row + 1)
                    if str(ws.cell(r, 2).value or "").strip().lower().startswith("source")), None)
    if end_row is None:
        msg = f"{workbook_path.name}: sheet '{ws.title}': no 'Source:' line in column B, so the end of the indicator table cannot be found"
        raise ValueError(msg)
    labels = {r: normalise_label(ws.cell(r, 2).value) for r in range(1, end_row)}
    resolved_rows: Dict[str, Tuple[int, str]] = {}
    for series_id, row_2026, expected, unit in INDICATORS:
        rows = [r for r, label in labels.items() if label == expected]
        if len(rows) != 1:
            found = f"rows {', '.join(map(str, rows))}" if rows else "none"
            msg = (f"{workbook_path.name}: sheet '{ws.title}': indicator {series_id} needs exactly one row labelled "
                   f"'{expected}' above the 'Source:' line (row {row_2026} in the 2026-10 workbook); found {found}")
            raise ValueError(msg)
        resolved_rows[series_id] = (rows[0], unit)

    # 2b. A formula whose result was never stored reads as blank; refuse it instead of dropping values
    formulas = openpyxl.load_workbook(workbook_path, read_only=True)
    try:
        needed = {row: series_id for series_id, (row, _) in resolved_rows.items()}
        first, last = min(needed), max(needed)
        for row, values in enumerate(formulas[ws.title].iter_rows(min_row=first, max_row=last, values_only=True), start=first):
            if row not in needed:
                continue
            for col in col_dates:
                f = values[col - 1] if col - 1 < len(values) else None
                if isinstance(f, str) and f.startswith("=") and ws.cell(row, col).value is None:
                    msg = (f"{workbook_path.name}: sheet '{ws.title}': cell {get_column_letter(col)}{row} ({needed[row]}) holds a "
                           f"formula with no stored result; recalculate and save the workbook in Excel before parsing")
                    raise ValueError(msg)
    finally:
        formulas.close()

    # 3. Extract observations
    observations: List[Dict] = []

    for series_id, (row_idx, unit) in resolved_rows.items():
        for col_idx, iso_date in col_dates.items():
            cell = ws.cell(row_idx, col_idx)
            # A blank or NO_VALUE cell is no value (counted); any other non-number stops the run
            numeric_val = source_io.read_value(cell.value, report, workbook_path, markers=NO_VALUE, sheet=SHEET,
                                               location=f"cell {cell.coordinate}", period=iso_date,
                                               series_id=series_id)
            if numeric_val is None:
                continue

            status_flag = "A"  # provisional; pipeline/flags.py sets the final flag from the flag policy

            observations.append({
                "date": iso_date,
                "series_id": series_id,
                "value": numeric_val,
                "unit": unit,
                "status_flag": status_flag,
                "realtime_start": REALTIME_START,
                "source_hash": source_hash,
            })

    df = pd.DataFrame(observations, columns=CANONICAL_COLUMNS)
    df = df.sort_values(by=["series_id", "date"]).reset_index(drop=True)

    # 4. Tier-1 validation
    # Final status flags from the flag policy (pipeline/metadata/flag_policy.yaml)
    df = flags.assign_flags(df, "money_banking_monthly")
    source_io.check_schema(df, workbook_path, "money_banking_monthly")
    report.finish()

    # 5. Output to Garden
    source_io.write_table(df, output_parquet_path)
    print(f"Saved {output_parquet_path} ({len(df):,} observations across {df['series_id'].nunique()} series).")
    print(f"Timeline: {df['date'].min()} to {df['date'].max()}.")

    return df


if __name__ == "__main__":
    parse_cbi_money_banking()
