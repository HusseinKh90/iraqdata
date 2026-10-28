"""
Exchange-rate sources.

parallel_rate_iraq comes from the CBI seed up to 2022-12-31 and from the Baghdad close in the
archived daily_fx.csv from 2023-01-01, with flag B on the first 2023 value; fx.official comes
from the archived CBI page-144 workbook for its whole history.
"""

import csv
import hashlib
import shutil

import pandas as pd
import pytest

import downloads
import parse_cbi_exchange_rates as fx
import vintages
from conftest import GARDEN_DIR, RAW_DIR, SEEDS_DIR
from fixtures import damage

pytestmark = pytest.mark.usefixtures("no_network")

SEED = SEEDS_DIR / "cbi_exchange_rate_historical.xlsx"
DAILY_FX = downloads.latest_version(RAW_DIR / "github" / "daily_fx.csv")
CBI_FX = downloads.latest_version(RAW_DIR / "cbi" / "fx" / "cbi_fx_usd.xlsx")
SPLICE = "2023-01-01"


def _hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def _parse(tmp_path, seed=SEED, cbi=CBI_FX, daily=DAILY_FX):
    out = tmp_path / "exchange_rates_daily.parquet"
    fx.parse_cbi_exchange_rates(seed_path=seed, cbi_workbook_path=cbi, daily_fx_path=daily,
                                output_parquet_path=out)
    return pd.read_parquet(out)


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    return _parse(tmp_path_factory.mktemp("fx"))


@pytest.fixture(scope="module")
def published():
    """The published table."""
    return pd.read_parquet(GARDEN_DIR / "exchange_rates_daily.parquet")


def _closes():
    with open(DAILY_FX, encoding="utf-8", newline="") as f:
        return {r["date"]: r["baghdad_close"] for r in csv.DictReader(f) if r["date"] >= SPLICE}


def _series(df, sid):
    return df[df["series_id"] == sid].set_index("date")


# ---------------------------------------------------------------- parallel rate

def test_parallel_up_to_2022_is_the_cbi_seed(built, published):
    new, old = _series(built, "parallel_rate_iraq"), _series(published, "parallel_rate_iraq")
    new, old = new[new.index < SPLICE], old[old.index < SPLICE]
    assert new.index.equals(old.index)
    assert (new["value"] == old["value"]).all() and (new["source_hash"] == old["source_hash"]).all()
    assert (new["status_flag"] == old["status_flag"]).all()
    assert set(new["source_hash"]) == {_hash(SEED)}


def test_parallel_from_2023_is_the_baghdad_close(built):
    rows = _series(built, "parallel_rate_iraq")
    rows = rows[rows.index >= SPLICE]
    expected = {d: float(c) for d, c in _closes().items() if c.strip()}
    assert dict(zip(rows.index, rows["value"])) == expected
    assert set(rows["source_hash"]) == {_hash(DAILY_FX)}


def test_days_without_a_close_are_omitted(built):
    blank = [d for d, c in _closes().items() if not c.strip()]
    assert blank, "today's file has days without a Baghdad close (2023-08-06 among them)"
    assert not built[(built.series_id == "parallel_rate_iraq") & built.date.isin(blank)].shape[0]


def test_exactly_one_break_flag_at_the_splice(built):
    flagged = built[built["status_flag"] == "B"]
    first = min(d for d, c in _closes().items() if c.strip())
    assert list(zip(flagged["series_id"], flagged["date"])) == [("parallel_rate_iraq", first)]


@pytest.mark.parametrize("change, words", [
    (lambda rows: [{k: v for k, v in r.items() if k != "baghdad_close"} for r in rows], ["baghdad_close"]),
    (lambda rows: rows + [rows[-1]], ["twice"]),
    (lambda rows: [dict(rows[0], date="2023-02-30")] + rows[1:], ["2023-02-30"]),
    (lambda rows: [dict(rows[0], baghdad_close="abc")] + rows[1:], ["abc"]),
    (lambda rows: [dict(rows[0], baghdad_close="inf")] + rows[1:], ["inf"]),
], ids=["missing_column", "duplicate_date", "impossible_date", "non_numeric_close", "infinite_close"])
def test_malformed_daily_fx_stops_the_run(change, words, tmp_path):
    with open(DAILY_FX, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    rows = change(rows)
    bad = tmp_path / DAILY_FX.name
    with open(bad, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    out = tmp_path / "out"
    out.mkdir()
    with pytest.raises(ValueError) as info:
        _parse(out, daily=bad)
    message = str(info.value)
    assert message.startswith(DAILY_FX.name), message
    assert all(w in message for w in words), message


# ---------------------------------------------------------------- official rate

def _rows(path):
    import openpyxl

    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Daily Price basis"]
    return list(ws.iter_rows(min_row=1, values_only=True))


def _header_row(rows, year):
    return next(i for i, r in enumerate(rows, 1) if r and str(r[0] or "").startswith(f"{year}/"))


def _cell(path, year, month, day):
    """The coordinate of one day in the CBI workbook's calendar blocks (column A = day, B = January)."""
    rows = _rows(path)
    header = _header_row(rows, year)
    row = next(i for i in range(header + 1, header + 35) if rows[i - 1][0] == day)
    return f"{chr(ord('A') + month)}{row}"


def test_official_comes_from_the_cbi_workbook(built, published):
    new, old = _series(built, "fx.official"), _series(published, "fx.official")
    assert set(new["source_hash"]) == {_hash(CBI_FX)}
    assert new.index.equals(old.index) and (new["value"] == old["value"]).all()
    assert new.index.max() > "2026-10-20"
    assert new.at["2026-10-06", "value"] == 1300.0 and new.at["2026-10-07", "value"] == 1500.0
    assert (new["status_flag"] == "A").all()


def test_merging_a_rebuild_into_the_published_table_changes_nothing(built, published):
    merged = vintages.merge_vintages(published, built, "2099-01-01", "exchange_rates_daily")
    assert not merged.changed and merged.table.equals(published)


def test_a_cbi_correction_to_a_past_day_is_a_revision(tmp_path, published):
    copy = tmp_path / CBI_FX.name
    shutil.copy2(CBI_FX, copy)
    cell = _cell(copy, 2010, 5, 3)
    import openpyxl
    old = openpyxl.load_workbook(copy, data_only=True)["Daily Price basis"][cell].value
    damage.edit_cell(copy, "Daily Price basis", cell, old + 5)
    out = tmp_path / "out"
    out.mkdir()
    res = vintages.merge_vintages(published, _parse(out, cbi=copy), "2099-01-01", "exchange_rates_daily")
    official = [r for r in res.revised if r["series_id"] == "fx.official"]
    assert [(r["date"], r["old_value"], r["new_value"]) for r in official] == [("2010-05-03", old, old + 5)]


def _damaged_cbi(tmp_path, edit):
    copy = tmp_path / CBI_FX.name
    shutil.copy2(CBI_FX, copy)
    edit(copy)
    out = tmp_path / "out"
    out.mkdir()
    with pytest.raises(ValueError) as info:
        _parse(out, cbi=copy)
    return str(info.value)


def test_cbi_workbook_without_the_sheet_stops(tmp_path):
    message = _damaged_cbi(tmp_path, lambda p: damage.remove_sheet(p, "Daily Price basis"))
    assert message.startswith(CBI_FX.name) and "Daily Price basis" in message and "Monthly price basis" in message


def test_cbi_missing_year_block_stops(tmp_path):
    message = _damaged_cbi(tmp_path, lambda p: damage.edit_cell(
        p, "Daily Price basis", f"A{_header_row(_rows(p), 2015)}", "Year 15 / Date"))
    assert message.startswith(CBI_FX.name) and "Daily Price basis" in message and "2015" in message


def test_cbi_formula_without_stored_result_stops(tmp_path):
    message = _damaged_cbi(tmp_path, damage.strip_formula_cache)
    assert message.startswith(CBI_FX.name) and "formula" in message
