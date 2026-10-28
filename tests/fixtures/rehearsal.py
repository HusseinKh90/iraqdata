"""
tests.fixtures.rehearsal
========================
Build a "next release" of each parser's input inside a sandbox
copy of pipeline/raw and pipeline/seeds. Nothing here writes under the real pipeline/.

Each builder takes (raw_root, seeds_root, date) and returns an Expectation:

- new:       (series_id, period) that must appear with realtime_start == date
- revised:   (series_id, period, new_value) that must hold the new value with the new vintage
- also:      other (series_id, period) keys the release legitimately changes (derived series
             such as BIS net position, every series of the new period)
- vintage:   the --vintage to pass (seed sources, whose files carry no date), else None

Parser-specific workarounds:
- cosit_gdp has no builder: its values are typed into the parser, so the test edits the fresh
  build through run_all's `_fresh_hook`.
- oil_monthly: the dated INTL zip holds only the Iraq line the parser reads (the full bulk file
  is 129 MB uncompressed).
- Seed workbooks (fiscal_seed) are overwritten in the sandbox copy and run with --vintage,
  because seeds carry no dated names. cbi_exchange_rates reads dated raw files.
- openpyxl writes floats with "%.16g", which drops the 17th significant digit of values Excel
  stored in full; a plain resave would then look like dozens of revisions. Workbooks are saved
  through `_save`, which writes floats with repr (full precision), as Excel does.
- The "omitted period" scenario is applied to the fresh build through `_fresh_hook` for every
  parser, since the merge rule it tests is the same for all tables.
"""

import contextlib
import io
import json
import re
import shutil
import warnings
import zipfile
from pathlib import Path
from typing import Callable, Dict, List, NamedTuple, Optional, Tuple

import openpyxl
import pandas as pd


class Expectation(NamedTuple):
    new: Tuple[str, str]
    revised: Tuple[str, str, float]
    also: List[Tuple[str, str]]
    vintage: Optional[str] = None


def _load(path: Path) -> openpyxl.Workbook:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=True)


@contextlib.contextmanager
def _full_precision():
    import openpyxl.cell._writer as writer

    original = writer.safe_string

    def safe_string(value):
        if isinstance(value, float) and value == value and abs(value) != float("inf"):
            return repr(value)
        return original(value)

    writer.safe_string = safe_string
    try:
        yield
    finally:
        writer.safe_string = original


def _save(wb: openpyxl.Workbook, path: Path) -> None:
    with _full_precision():
        wb.save(path)


def _dated(path: Path, date: str) -> Path:
    return path.with_name(f"{path.stem}_{date}{path.suffix}")


def _num(v) -> Optional[float]:
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


# ===================================================================== CSV / JSON / ZIP sources


def bis_lbs(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    bis = raw_root / "bis"
    revised_value = None
    for src in sorted(bis.glob("bis_lbs_*.csv")):
        if src.stem.count("_") > 3:  # already a dated copy
            continue
        df = pd.read_csv(src, dtype=str, keep_default_na=False)
        last = df.iloc[[-1]].copy()
        last["TIME_PERIOD"] = "2026-Q2"
        last["OBS_VALUE"] = f"{float(last['OBS_VALUE'].iloc[0]) * 1.01:.3f}"
        df = pd.concat([df, last], ignore_index=True)
        if src.name == "bis_lbs_claims_total.csv":
            i = df.index[df["TIME_PERIOD"] == "2025-Q4"][0]
            revised_value = round(float(df.at[i, "OBS_VALUE"]) + 100.0, 3)
            df.at[i, "OBS_VALUE"] = f"{revised_value:.3f}"
        df.to_csv(_dated(src, date), index=False)
    series = ["bis.lbs.claims_total", "bis.lbs.liabilities_total", "bis.lbs.liabilities_banks",
              "bis.lbs.liabilities_nonbank", "bis.lbs.claims_nonbank", "bis.lbs.net_position"]
    return Expectation(("bis.lbs.claims_total", "2026-Q2"),
                       ("bis.lbs.claims_total", "2025-Q4", revised_value),
                       [(s, "2026-Q2") for s in series] + [("bis.lbs.net_position", "2025-Q4")])


def oecd_dac(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    oecd = raw_root / "oecd"
    revised_value = None
    for name in ("wb_oda_net.json", "wb_oda_grants.json"):
        src = oecd / name
        data = json.loads(src.read_text(encoding="utf-8"))
        by_year = {item["date"]: item for item in data[1]}
        by_year["2024"]["value"] = by_year["2023"]["value"] * 0.9
        if name == "wb_oda_net.json":
            by_year["2015"]["value"] = by_year["2015"]["value"] + 50_000_000.0
            revised_value = round(by_year["2015"]["value"] / 1_000_000.0, 3)
        _dated(src, date).write_text(json.dumps(data), encoding="utf-8")
    series = ["oecd.dac.oda_total_net", "oecd.dac.grants_excl_tc"]
    return Expectation(("oecd.dac.oda_total_net", "2024"), ("oecd.dac.oda_total_net", "2015", revised_value),
                       [(s, "2024") for s in series] + [(s, "2015") for s in series])


def oil_daily(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    src = raw_root / "fred" / "DCOILBRENTEU.csv"
    df = pd.read_csv(src, dtype=str, keep_default_na=False)
    last_day = pd.Timestamp(df["observation_date"].iloc[-1])
    new_day = (last_day + pd.Timedelta(days=1)).date().isoformat()
    df = pd.concat([df, pd.DataFrame({"observation_date": [new_day], "DCOILBRENTEU": ["100.00"]})], ignore_index=True)
    i = df.index[df["observation_date"] == "2026-10-01"][0]
    revised_value = round(float(df.at[i, "DCOILBRENTEU"]) + 1.0, 2)
    df.at[i, "DCOILBRENTEU"] = f"{revised_value:.2f}"
    df.to_csv(_dated(src, date), index=False)
    return Expectation(("oil_price_brent", new_day), ("oil_price_brent", "2026-10-01", revised_value), [])


def _eia_iraq_line(zip_path: Path) -> dict:
    with zipfile.ZipFile(zip_path) as z, z.open("INTL.txt") as f:
        for line in f:
            if b"INTL.57-1-IRQ-TBPD.M" in line:
                return json.loads(line.decode("utf-8"))
    raise ValueError("Iraq line not found")


def _write_eia(obj: dict, target: Path) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("INTL.txt", json.dumps(obj) + "\r\n")
    target.write_bytes(buf.getvalue())


def oil_monthly(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    src = raw_root / "eia" / "INTL.zip"
    obj = _eia_iraq_line(src)
    data = obj["data"]  # newest first: [["202606", 4321], ...]
    last = data[0][0]
    nxt = (pd.Period(f"{last[:4]}-{last[4:]}", freq="M") + 1).strftime("%Y%m")
    data.insert(0, [nxt, data[0][1]])
    for row in data:
        if row[0] == "202601":
            row[1] = float(row[1]) + 10.0
            revised_value = round(float(row[1]), 1)
    obj["end"] = nxt
    _write_eia(obj, _dated(src, date))
    return Expectation(("crude_production_secondary", f"{nxt[:4]}-{nxt[4:]}"),
                       ("crude_production_secondary", "2026-01", revised_value), [])


# ===================================================================== CBI workbooks


def cbi_bop(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    src = raw_root / "cbi" / "bop" / "bop_2026.xlsx"
    wb = _load(src)
    q1 = next(ws for ws in wb.worksheets if ws.title.strip().upper().endswith("Q1"))
    q2 = wb.copy_worksheet(q1)
    q2.title = "BOP Q2"
    for row in q2.iter_rows(max_row=15):  # the table's own title must name the new quarter, as CBI's does
        for cell in row:
            if isinstance(cell.value, str) and "balance of payments" in cell.value.lower() and "q1" in cell.value.lower():
                cell.value = re.sub(r"(?i)q1", lambda m: "Q2" if m.group() == "Q1" else "q2", cell.value)
    revised_value = None
    for row in q1.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "services" in cell.value.lower() and "account" in cell.value.lower():
                target = q1.cell(cell.row, cell.column + 1)
                if _num(target.value) is not None:
                    revised_value = round(_num(target.value) + 10.0, 1)
                    target.value = revised_value
                    break
        if revised_value is not None:
            break
    _save(wb, _dated(src, date))
    series = ["bop.current_account", "bop.trade_balance", "trade.exports_total", "trade.imports_total",
              "trade.imports_government", "trade.imports_private", "bop.services_net", "bop.income_net",
              "bop.transfers_net", "bop.capital_account", "bop.financial_account", "bop.reserve_assets_change",
              "trade.exports_crude_oil", "trade.exports_oil_products", "trade.exports_nonoil"]
    return Expectation(("bop.current_account", "2026-Q2"), ("bop.services_net", "2026-Q1", revised_value),
                       [(s, "2026-Q2") for s in series])


MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_MONTH_NUM = {m.lower(): i + 1 for i, m in enumerate(MONTHS)}


def _kfi_timeline(ws) -> List[Tuple[int, int, int]]:
    """(column, year, month) for every monthly column, as the parsers read them."""
    out, year = [], None
    for col in range(3, ws.max_column + 1):
        r1 = ws.cell(1, col).value
        if r1 is not None:
            s = str(r1).strip()
            if "week ending" in s.lower():
                continue
            if s.isdigit() and len(s) == 4:
                year = int(s)
        r2 = ws.cell(2, col).value
        if r2 is None or year is None:
            continue
        key = str(r2).strip().lower().rstrip(".")[:3]
        if key in _MONTH_NUM:
            out.append((col, year, _MONTH_NUM[key]))
    return out


def _kfi(raw_root: Path, date: str, revise_row: int, revise_period: str) -> Tuple[str, float]:
    """Append next month's column (each row repeats its value from the last 3 months, so ended
    series stay ended) and change one cell."""
    src = raw_root / "cbi" / "key_financial_indicators.xlsx"
    wb = _load(src)
    ws = wb.active
    timeline = _kfi_timeline(ws)
    last_col, year, month = timeline[-1]
    recent = [c for c, _, _ in timeline[-3:]]
    ny, nm = (year + 1, 1) if month == 12 else (year, month + 1)
    rev_col = next(c for c, y, m in timeline if f"{y}-{m:02d}" == revise_period)
    revised_value = round(_num(ws.cell(revise_row, rev_col).value) + 1.0, 4)
    ws.cell(revise_row, rev_col).value = revised_value
    new_col = last_col + 1
    ws.insert_cols(new_col)
    ws.cell(1, new_col).value = str(ny) if ny != year else None
    ws.cell(2, new_col).value = MONTHS[nm - 1]
    for r in range(3, ws.max_row + 1):
        for c in reversed(recent):
            v = _num(ws.cell(r, c).value)
            if v is not None:
                ws.cell(r, new_col).value = v
                break
    _save(wb, raw_root / "cbi" / f"cbi_key_financial_indicators_{date}.xlsx")
    return f"{ny}-{nm:02d}", revised_value


def cbi_money_banking(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    import parse_cbi_money_banking as p

    row = next(r for sid, r, _, _ in p.INDICATORS if sid == "money.m0")
    new_period, revised_value = _kfi(raw_root, date, row, "2026-06")
    return Expectation(("money.m0", new_period), ("money.m0", "2026-06", revised_value), [])


def cbi_cpi(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    import parse_cbi_cpi as p

    row = next(r for sid, r, _, _ in p.CPI_NATIVE_ROWS if sid == "cpi.headline_2022")
    new_period, revised_value = _kfi(raw_root, date, row, "2026-06")
    return Expectation(("cpi.headline_2022", new_period), ("cpi.headline_2022", "2026-06", revised_value),
                       [("cpi_headline", "2026-06")])


def cbi_fiscal_bulletin(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    import parse_cbi_fiscal_bulletin as p

    src = raw_root / "cbi" / "bulletin_latest.xlsx"
    wb = _load(src)
    ws = wb["6-1"]
    year, dec_2025, header_2026 = None, None, None
    for r in range(9, ws.max_row + 1):
        s = str(ws.cell(r, 1).value or "").strip()
        if s.isdigit() and len(s) == 4:
            year = s
            if s == "2026":
                header_2026 = r
        elif year == "2025" and "dec" in s.lower():
            dec_2025 = r
    col_rev = p.FISCAL_SERIES_COLS["fiscal.revenue_total"]
    revised_value = float(ws.cell(dec_2025, col_rev).value) + 1000.0
    ws.cell(dec_2025, col_rev).value = revised_value
    ws.insert_rows(header_2026 + 1)
    ws.cell(header_2026 + 1, 1).value = "Dec"
    for col in p.FISCAL_SERIES_COLS.values():
        ws.cell(header_2026 + 1, col).value = ws.cell(dec_2025, col).value
    _save(wb, raw_root / "cbi" / f"cbi_bulletin_{date}.xlsx")
    series = list(p.FISCAL_SERIES_COLS)
    return Expectation(("fiscal.revenue_total", "2026"), ("fiscal.revenue_total", "2025", revised_value),
                       [(s, "2026") for s in series])


# ===================================================================== seeds (run with --vintage)


def fiscal_seed(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    path = seeds_root / "iraq_fiscal_historical.xlsx"
    df = pd.read_excel(path)
    new = df.iloc[[-1]].copy()
    new["year"] = 2026
    df = pd.concat([df, new], ignore_index=True)
    i = df.index[df["year"] == 2010][0]
    df.at[i, "revenue_actual_total"] = float(df.at[i, "revenue_actual_total"]) + 1000.0
    revised_value = float(df.at[i, "revenue_actual_total"])
    with _full_precision():
        df.to_excel(path, index=False)
    import parse_fiscal_seed as p
    return Expectation(("fiscal.revenue_total", "2026"), ("fiscal.revenue_total", "2010", revised_value),
                       [(s, "2026") for s in p.FISCAL_SERIES_MAP.values()], vintage=date)


def _cbi_fx_cell(ws, year: int, month: int, day: int):
    """The cell of one day in a calendar-block sheet (column A = day, B = January)."""
    header = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith(f"{year}/"))
    row = next(r for r in range(header + 1, header + 35) if _num(ws.cell(r, 1).value) == day)
    return ws.cell(row, month + 1)


def cbi_exchange_rates(raw_root: Path, seeds_root: Path, date: str) -> Expectation:
    """
    The next daily_fx.csv adds a day of parallel_rate_iraq; the next CBI workbook
    revises an old fx.official day (published with the 2026-10-22 vintage). Both are written as
    new dated files, as the fetchers do.
    """
    import csv

    import downloads

    daily = downloads.latest_version(raw_root / "github" / "daily_fx.csv")
    with open(daily, encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        fields, rows = reader.fieldnames, list(reader)
    last = pd.Timestamp(rows[-1]["date"])
    new_day = (last + pd.Timedelta(days=1)).strftime("%Y-%m-%d")
    rows.append(dict(rows[-1], date=new_day, baghdad_close="1600.0"))
    with open(_dated(raw_root / "github" / "daily_fx.csv", date), "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    workbook = downloads.latest_version(raw_root / "cbi" / "fx" / "cbi_fx_usd.xlsx")
    wb = _load(workbook)
    cell = _cbi_fx_cell(wb["Daily Price basis"], 2010, 5, 3)
    revised_value = float(_num(cell.value)) + 5.0
    cell.value = revised_value
    _save(wb, _dated(raw_root / "cbi" / "fx" / "cbi_fx_usd.xlsx", date))
    return Expectation(("parallel_rate_iraq", new_day), ("fx.official", "2010-05-03", revised_value), [])


BUILDERS: Dict[str, Callable[[Path, Path, str], Expectation]] = {
    "bis_lbs": bis_lbs,
    "cbi_bop": cbi_bop,
    "cbi_cpi": cbi_cpi,
    "cbi_exchange_rates": cbi_exchange_rates,
    "fiscal_seed": fiscal_seed,
    "cbi_fiscal_bulletin": cbi_fiscal_bulletin,
    "cbi_money_banking": cbi_money_banking,
    "oecd_dac": oecd_dac,
    "oil_daily": oil_daily,
    "oil_monthly": oil_monthly,
}


def identical_copy(parser: str, raw_root: Path, seeds_root: Path, date: str) -> Optional[str]:
    """Byte-identical dated copies of the parser's raw inputs. Returns --vintage for seeds."""
    copies = {
        "bis_lbs": [raw_root / "bis" / n for n in (
            "bis_lbs_claims_total.csv", "bis_lbs_liabilities_total.csv", "bis_lbs_liabilities_banks.csv",
            "bis_lbs_liabilities_nonbank.csv", "bis_lbs_claims_nonbank.csv")],
        "cbi_bop": [raw_root / "cbi" / "bop" / "bop_2026.xlsx"],
        "oecd_dac": [raw_root / "oecd" / "wb_oda_net.json", raw_root / "oecd" / "wb_oda_grants.json"],
        "oil_daily": [raw_root / "fred" / "DCOILBRENTEU.csv"],
        "oil_monthly": [raw_root / "eia" / "INTL.zip"],
    }
    if parser in copies:
        for src in copies[parser]:
            shutil.copyfile(src, _dated(src, date))
        return None
    if parser in ("cbi_cpi", "cbi_money_banking"):
        shutil.copyfile(raw_root / "cbi" / "key_financial_indicators.xlsx",
                        raw_root / "cbi" / f"cbi_key_financial_indicators_{date}.xlsx")
        return None
    if parser == "cbi_fiscal_bulletin":
        shutil.copyfile(raw_root / "cbi" / "bulletin_latest.xlsx", raw_root / "cbi" / f"cbi_bulletin_{date}.xlsx")
        return None
    if parser == "cbi_exchange_rates":  # the two raw sources (the seed is unchanged)
        import downloads
        for base in (raw_root / "github" / "daily_fx.csv", raw_root / "cbi" / "fx" / "cbi_fx_usd.xlsx"):
            shutil.copyfile(downloads.latest_version(base), _dated(base, date))
        return None
    return date  # seeds and typed-in values: unchanged input, explicit vintage
