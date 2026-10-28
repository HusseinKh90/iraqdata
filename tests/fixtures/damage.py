"""
Damaged-copy helpers.

Every helper works on copies under a test's tmp_path; nothing here touches pipeline/raw/,
pipeline/seeds/ or pipeline/data/. Workbooks are loaded with data_only=True (cached values) before
an edit, so untouched cells keep their values; only strip_formula_cache keeps the formulas, which
is what removes their stored results.
"""

import io
import shutil
import warnings
from pathlib import Path
from typing import List, Tuple

import openpyxl

import run_parsers

REPO_ROOT = Path(__file__).resolve().parents[2]
RAW = REPO_ROOT / "pipeline" / "raw"
SEEDS = REPO_ROOT / "pipeline" / "seeds"


def _load(path: Path, data_only: bool = True) -> openpyxl.Workbook:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, data_only=data_only)


def copy_inputs(parser: str, tmp_path: Path) -> Tuple[Path, Path]:
    """Copy the inputs of `parser` (and of the parsers it runs after) into tmp_path/raw and tmp_path/seeds."""
    raw_root, seeds_root = tmp_path / "raw", tmp_path / "seeds"
    names = list(run_parsers.BY_NAME[parser].after) + [parser]
    for name in names:
        for glob in run_parsers.BY_NAME[name].inputs:
            for src in REPO_ROOT.glob(glob):
                src = src.resolve()
                if src.is_relative_to(RAW.resolve()):
                    dest = raw_root / src.relative_to(RAW.resolve())
                else:
                    dest = seeds_root / src.relative_to(SEEDS.resolve())
                dest.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, dest)
    return raw_root, seeds_root


def run_parser(parser: str, raw_root: Path, seeds_root: Path, out: Path) -> Path:
    """Run `parser` (after its dependencies) offline-style on the given roots; returns `out`."""
    (out / "seeds").mkdir(parents=True, exist_ok=True)
    for name in list(run_parsers.BY_NAME[parser].after) + [parser]:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            run_parsers.BY_NAME[name].run(raw_root, seeds_root, out, out / "seeds")
    return out


def edit_cell(path: Path, sheet: str, cell: str, value) -> None:
    wb = _load(path)
    wb[sheet][cell] = value
    wb.save(path)


def insert_row(path: Path, sheet: str, row: int) -> None:
    wb = _load(path)
    wb[sheet].insert_rows(row)
    wb.save(path)


def swap_rows(path: Path, sheet: str, a: int, b: int) -> None:
    """Swap every cell of rows a and b (labels and values)."""
    wb = _load(path)
    ws = wb[sheet]
    for col in range(1, ws.max_column + 1):
        ca, cb = ws.cell(a, col), ws.cell(b, col)
        ca.value, cb.value = cb.value, ca.value
    wb.save(path)


def remove_sheet(path: Path, sheet: str) -> None:
    wb = _load(path)
    del wb[sheet]
    wb.save(path)


def strip_formula_cache(path: Path) -> None:
    """Re-save with formulas kept: openpyxl writes no stored results, as an unrecalculated file would."""
    _load(path, data_only=False).save(path)


def workbook_bytes(sheets: List[str]) -> bytes:
    """A minimal valid .xlsx with the given sheet names."""
    wb = openpyxl.Workbook()
    wb.active.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
