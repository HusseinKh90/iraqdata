"""
Status flag policy.

Unit tests use a synthetic policy and catalog; the repository-wide checks at the end run
`assign_flags` on real parser output.
"""

import datetime
import re

import numpy as np
import pandas as pd
import pytest

import flags

POLICY = {
    "precedence": ["M", "B", "E", "U", "P", "A"],
    "class_flags": {"observed": "A", "derived": "inherit", "modelled": "E", "transcribed": "U", "unknown": "U"},
    "default_windows": {"daily": 0, "monthly": 12, "quarterly": 2, "annual": 2},
    "series": {"pub.series": {"revision_window": "publisher", "rationale": "test"}},
}


def _entry(sid, freq, *segments):
    return {"series_id": sid, "frequency": freq, "origin": list(segments)}


def _seg(start, end, cls, **extra):
    return {"from": start, "to": end, "class": cls, "official": True, "method": "m", "source": "s", **extra}


CATALOG = {
    "obs.q": _entry("obs.q", "quarterly", _seg("2000-Q1", "open", "observed")),
    "mixed.q": _entry("mixed.q", "quarterly", _seg("2000-Q1", "2000-Q2", "observed"),
                      _seg("2000-Q3", "open", "transcribed", record="r.csv")),
    "total.q": _entry("total.q", "quarterly", _seg("2000-Q1", "open", "transcribed", record="r.csv")),
    "oil.q": _entry("oil.q", "quarterly", _seg("2000-Q1", "open", "transcribed", record="r.csv")),
    "nonoil.q": _entry("nonoil.q", "quarterly",
                       _seg("2000-Q1", "open", "derived", inputs=["total.q", "oil.q"], record="r.csv")),
    "model.m": _entry("model.m", "monthly", _seg("2000-01", "open", "modelled")),
    "daily.d": _entry("daily.d", "daily", _seg("2000-01-01", "open", "observed")),
    "pub.series": _entry("pub.series", "annual", _seg(2000, "open", "observed")),
}


def _frame(rows):
    return pd.DataFrame([{"date": d, "series_id": s, "value": v, "unit": "X", "status_flag": f,
                          "realtime_start": "2026-10-22", "source_hash": "0" * 16} for s, d, v, f in rows])


def _flags(df):
    out = flags.assign_flags(df, "t", policy=POLICY, catalog=CATALOG)
    return dict(zip(zip(out["series_id"], out["date"]), out["status_flag"]))


QUARTERS = ["2000-Q1", "2000-Q2", "2000-Q3", "2000-Q4"]


def test_observed_outside_window_is_A_inside_is_P():
    got = _flags(_frame([("obs.q", q, 1.0, "A") for q in QUARTERS]))
    assert [got[("obs.q", q)] for q in QUARTERS] == ["A", "A", "P", "P"]


def test_transcribed_outranks_window():
    got = _flags(_frame([("mixed.q", q, 1.0, "A") for q in QUARTERS]))
    assert [got[("mixed.q", q)] for q in QUARTERS] == ["A", "A", "U", "U"]


def test_break_outranks_transcribed_and_nan_is_M():
    got = _flags(_frame([("mixed.q", "2000-Q1", 1.0, "A"), ("mixed.q", "2000-Q2", np.nan, "A"),
                         ("mixed.q", "2000-Q3", 1.0, "B"), ("mixed.q", "2000-Q4", 1.0, "A")]))
    assert got[("mixed.q", "2000-Q2")] == "M"
    assert got[("mixed.q", "2000-Q3")] == "B"
    assert got[("mixed.q", "2000-Q4")] == "U"


def test_modelled_is_E():
    got = _flags(_frame([("model.m", "2000-01", 1.0, "A")]))
    assert got[("model.m", "2000-01")] == "E"


def test_window_zero_gives_no_P():
    got = _flags(_frame([("daily.d", "2000-01-0%d" % i, 1.0, "A") for i in range(1, 4)]))
    assert set(got.values()) == {"A"}


def test_publisher_window_keeps_only_the_sources_P():
    got = _flags(_frame([("pub.series", "2000", 1.0, "A"), ("pub.series", "2001", 1.0, "A"),
                         ("pub.series", "2002", 1.0, "P")]))
    assert got == {("pub.series", "2000"): "A", ("pub.series", "2001"): "A", ("pub.series", "2002"): "P"}


def test_parser_P_is_ignored_when_the_window_decides():
    got = _flags(_frame([("obs.q", "2000-Q1", 1.0, "P"), ("obs.q", "2000-Q2", 1.0, "A"),
                         ("obs.q", "2000-Q3", 1.0, "A")]))
    assert got[("obs.q", "2000-Q1")] == "A"


def test_window_counts_on_each_series_own_latest_period():
    rows = [("obs.q", q, 1.0, "A") for q in QUARTERS] + [("total.q", "2000-Q1", 1.0, "A")]
    got = _flags(_frame(rows))
    assert got[("obs.q", "2000-Q3")] == "P"  # obs.q's latest two, not the table's


def test_derived_inherits_from_its_inputs():
    rows = []
    for q in QUARTERS:
        rows += [("total.q", q, 10.0, "A"), ("oil.q", q, 4.0, "A"), ("nonoil.q", q, 6.0, "A")]
    got = _flags(_frame(rows))
    assert {got[("nonoil.q", q)] for q in QUARTERS} == {"U"}


def test_unknown_series_raises():
    with pytest.raises(ValueError, match="nope"):
        _flags(_frame([("nope", "2000-Q1", 1.0, "A")]))


def test_period_without_segment_raises():
    with pytest.raises(ValueError, match="obs.q"):
        _flags(_frame([("obs.q", "1999-Q4", 1.0, "A")]))


def test_only_status_flag_changes_and_input_is_untouched():
    df = _frame([("mixed.q", q, 1.0, "A") for q in QUARTERS])
    before = df.copy()
    out = flags.assign_flags(df, "t", policy=POLICY, catalog=CATALOG)
    assert df.equals(before)
    assert out.drop(columns="status_flag").equals(before.drop(columns="status_flag"))


def test_the_clock_has_no_effect(monkeypatch):
    df = _frame([("obs.q", q, 1.0, "A") for q in QUARTERS])
    first = _flags(df)

    class Later(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime.datetime(2040, 1, 1, tzinfo=tz)
    monkeypatch.setattr(datetime, "datetime", Later)
    monkeypatch.setattr(pd.Timestamp, "today", classmethod(lambda cls, *a, **k: pd.Timestamp("2040-01-01")))
    assert _flags(df) == first


def test_repository_policy_is_well_formed():
    policy = flags.load_policy()
    assert policy["precedence"] == ["M", "B", "E", "U", "P", "A"]
    assert set(policy["class_flags"]) == set(flags.ORIGIN_CLASSES)
    catalog = flags.load_catalog()
    for sid, rule in (policy.get("series") or {}).items():
        assert sid in catalog, sid
        assert rule["revision_window"] == "publisher" or int(rule["revision_window"]) >= 0, sid
        assert str(rule.get("rationale", "")).strip(), sid


def test_repository_breaks_are_valid_and_only_confirmed_ones_apply():
    policy = flags.load_policy()
    breaks = policy["breaks"]
    assert breaks and {b["status"] for b in breaks} <= set(flags.BREAK_STATUSES)
    flags.confirmed_breaks(policy, flags.load_catalog())  # raises on an unknown series, a bad period or a repeat
    applied = {(b["series"], str(b["period"])) for b in breaks if b["status"] == "confirmed"}
    assert set(flags.confirmed_breaks(policy)) == applied


# =====================================================================
# Known breaks (a source's own footnote names the period a definition changed)
# =====================================================================

def _break(series="obs.q", when="2000-Q3", status="confirmed", **extra):
    return {"series": series, "period": when, "status": status, "source": "note 1", "note": "the definition changed", **extra}


def _flags_with(breaks, rows):
    out = flags.assign_flags(_frame(rows), "t", policy={**POLICY, "breaks": breaks}, catalog=CATALOG)
    return dict(zip(zip(out["series_id"], out["date"]), out["status_flag"]))


def test_a_confirmed_break_flags_only_its_period():
    got = _flags_with([_break()], [("obs.q", q, 1.0, "A") for q in QUARTERS])
    assert got[("obs.q", "2000-Q3")] == "B" and got[("obs.q", "2000-Q2")] == "A" and got[("obs.q", "2000-Q4")] == "P"
    # B outranks the provisional window that would otherwise give the latest periods P
    got = _flags_with([_break(when="2000-Q4")], [("obs.q", q, 1.0, "A") for q in QUARTERS])
    assert got[("obs.q", "2000-Q4")] == "B"


def test_a_proposed_break_changes_nothing():
    rows = [("obs.q", q, 1.0, "A") for q in QUARTERS]
    assert _flags_with([_break(status="proposed")], rows) == _flags_with([], rows)


def test_a_break_outranks_transcribed_and_passes_to_a_derived_series():
    rows = [(s, q, 1.0, "A") for s in ("total.q", "oil.q", "nonoil.q") for q in QUARTERS[:2]]
    got = _flags_with([_break(series="oil.q", when="2000-Q2")], rows)
    assert got[("oil.q", "2000-Q2")] == "B" and got[("oil.q", "2000-Q1")] == "U"
    assert got[("nonoil.q", "2000-Q2")] == "B"   # a derived row takes the strongest flag of its inputs


@pytest.mark.parametrize("bad, message", [
    (_break(series="nowhere.q"), "not in series_catalog"),
    (_break(when="2000-03"), "not a quarterly period"),
    (_break(status="maybe"), "status must be"),
    (_break(note=" "), "note must be"),
    (_break(source=""), "source must be"),
])
def test_a_malformed_break_is_refused_and_named(bad, message):
    with pytest.raises(ValueError, match=message) as info:
        flags.confirmed_breaks({"breaks": [bad]}, CATALOG)
    assert "breaks[0]" in str(info.value)


def test_a_break_listed_twice_is_refused():
    with pytest.raises(ValueError, match="listed twice"):
        flags.confirmed_breaks({"breaks": [_break(), _break(status="proposed")]}, CATALOG)


def test_the_kfi_proposals_name_real_footnotes_of_the_workbook():
    """Every proposed KFI break quotes the workbook's own note, so the maintainer can check it against the source."""
    import openpyxl
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(REPO_ROOT / "pipeline" / "raw" / "cbi" / "key_financial_indicators.xlsx",
                                    data_only=True).worksheets[0]
    squash = lambda text: re.sub(r"\s+", "", text).lower()  # noqa: E731  the workbook's spacing is irregular
    notes = squash(" ".join(str(ws.cell(r, 2).value) for r in range(118, 138) if ws.cell(r, 2).value))
    for entry in flags.load_policy()["breaks"]:
        assert squash(entry["note"].split(" (May be")[0]) in notes, entry["note"]


# =====================================================================
# Repository-wide: every flag in parser output comes from the policy
# =====================================================================

from conftest import REPO_ROOT  # noqa: E402

PARSERS = sorted((REPO_ROOT / "pipeline" / "parsers").glob("*.py"))
DATE_FLAG_RULE = re.compile(r'"P"\s+if\s+.*(>=|<=|>|<|==)\s*("?\d{4}|max_date)')
CONDITIONAL_FLAG = re.compile(r'flag\s*=\s*"[AP]"\s+if\b')


@pytest.mark.parametrize("path", PARSERS, ids=lambda p: p.name)
def test_no_date_literal_flag_rules(path):
    """No parser sets a flag from a calendar date or the latest period."""
    for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        assert not DATE_FLAG_RULE.search(line), f"{path.name}:{n}: {line.strip()}"
        assert not CONDITIONAL_FLAG.search(line), f"{path.name}:{n}: {line.strip()}"


def test_parser_outputs_are_policy_flags(sandbox_runs):
    """Every table a parser writes equals assign_flags of itself."""
    import run_parsers
    outputs = [(p, p.stem) for p in sorted(sandbox_runs.run1.glob("*.parquet"))]
    outputs += [(p, run_parsers.SNAPSHOT_TABLE[p.stem]) for p in sorted((sandbox_runs.run1 / "seeds").glob("*.parquet"))
                if p.stem in run_parsers.SNAPSHOT_TABLE]
    assert len(outputs) >= 10
    for path, table in outputs:
        df = pd.read_parquet(path)
        assert flags.assign_flags(df, table)["status_flag"].tolist() == df["status_flag"].tolist(), path.name


def test_rebuilt_flags_ignore_the_clock(sandbox_runs, tmp_path, monkeypatch):
    """Building again with the clock five years ahead gives the same flags."""
    import run_parsers

    class Later(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime.datetime(2031, 10, 25, tzinfo=tz)

        @classmethod
        def today(cls):
            return datetime.datetime(2031, 10, 25)
    monkeypatch.setattr(datetime, "datetime", Later)
    monkeypatch.setattr(pd.Timestamp, "today", classmethod(lambda cls, *a, **k: pd.Timestamp("2031-10-25")))
    run_parsers.run_all(out_dir=tmp_path, only=["oil_monthly", "cosit_gdp", "cbi_bop"])
    for name in ("oil_monthly", "gdp_quarterly", "bop_trade_quarterly"):
        later = pd.read_parquet(tmp_path / f"{name}.parquet")
        now = pd.read_parquet(sandbox_runs.run1 / f"{name}.parquet")
        assert later["status_flag"].tolist() == now["status_flag"].tolist(), name
