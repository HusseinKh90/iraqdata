"""
tests.test_kfi_timeline
=======================
Both parsers that read the CBI Key Financial Indicators workbook (cbi_cpi, cbi_money_banking) take
its month columns from one reader, pipeline/parsers/kfi_timeline.py. A header it does not
recognise, a gap or repeat in the months, or a missing year stops the run with the cell named;
before, such a column was skipped silently and its month disappeared from the output.
"""

from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

import kfi_timeline
from conftest import REPO_ROOT, block_network
from fixtures import damage
from source_io import ParseError

KFI = "key_financial_indicators.xlsx"
PATH = Path("kfi.xlsx")


def sheet(columns):
    """A worksheet whose columns 3.. hold (year header or None, month header) pairs."""
    ws = openpyxl.Workbook().active
    ws.title = "Sheet1"
    for i, (year, month) in enumerate(columns, start=3):
        ws.cell(1, i).value = year
        ws.cell(2, i).value = month
    return ws


# --------------------------------------------------------------------- month labels

@pytest.mark.parametrize("label, month", [
    ("Jan.", 1), ("Feb ", 2), ("March", 3), ("Apr", 4), ("May.", 5), ("June", 6), ("Jul", 7),
    ("Aug.", 8), ("Sept.", 9), ("Sep", 9), ("Sept", 9), ("sep.", 9), ("Oct.", 10), ("Nov", 11), ("December", 12),
])
def test_known_spellings(label, month):
    assert kfi_timeline.month_of(label) == month


@pytest.mark.parametrize("label", ["Augst", "Septembre", "", None, 7, "2026-10-01", "Q1", "Mon"])
def test_unknown_labels_are_not_months(label):
    assert kfi_timeline.month_of(label) is None


# --------------------------------------------------------------------- the timeline

def test_months_are_dated_from_the_year_header():
    ws = sheet([(2003, "Dec."), (2004, "Jan."), (None, "Feb"), (None, "March")])
    assert kfi_timeline.resolve(ws, PATH) == {4: "2004-01", 5: "2004-02", 6: "2004-03"}  # 2003-12 is before the sample


def test_year_may_be_text_float_or_arabic_indic():
    ws = sheet([("2004", "Dec."), ("٢٠٠٥", "Jan."), (2005.0, "Feb")])
    assert kfi_timeline.resolve(ws, PATH) == {3: "2004-12", 4: "2005-01", 5: "2005-02"}


def test_weekly_block_ends_the_monthly_timeline():
    ws = sheet([(2026, "Jul"), (None, "Aug"), ("The week ending ", "2026-10-01"), (None, None)])
    assert kfi_timeline.resolve(ws, PATH) == {3: "2026-07", 4: "2026-08"}


def test_trailing_empty_columns_are_ignored():
    ws = sheet([(2026, "Jul"), (None, "Aug"), (None, None), (None, None)])
    assert kfi_timeline.resolve(ws, PATH) == {3: "2026-07", 4: "2026-08"}


def _fails(columns, **kwargs):
    with pytest.raises(ParseError) as info:
        kfi_timeline.resolve(sheet(columns), PATH, **kwargs)
    return str(info.value)


def test_unrecognised_month_header_is_named():
    msg = _fails([(2026, "Jul"), (None, "Augst"), (None, "Sep")])
    assert "cell D2" in msg and "'Augst'" in msg and "sheet 'Sheet1'" in msg and "kfi.xlsx" in msg


def test_blank_month_header_between_months_is_named():
    msg = _fails([(2026, "Jul"), (None, None), (None, "Sep")])
    assert "cell D2" in msg and "None" in msg


def test_a_skipped_month_is_named():
    msg = _fails([(2026, "Jul"), (None, "Sep")])
    assert "2026-08" in msg and "2026-09" in msg and "cell D2" in msg


def test_a_repeated_month_is_named():
    msg = _fails([(2026, "Jul"), (None, "Jul")])
    assert "2026-08" in msg and "2026-07" in msg


def test_a_forgotten_year_header_is_named():
    """Without the new year's header, January would repeat the old year and collide with it."""
    msg = _fails([(2025, "Nov"), (None, "Dec"), (None, "Jan")])
    assert "2026-01" in msg and "2025-01" in msg and "cell E2" in msg


def test_month_before_any_year_is_named():
    assert "year" in _fails([(None, "Jan")])


def test_unexpected_text_in_a_year_cell_is_named():
    msg = _fails([(2026, "Jan"), ("Revised", "Feb")])
    assert "cell D1" in msg and "'Revised'" in msg


def test_a_sheet_without_months_is_refused():
    assert "monthly" in _fails([(2003, "Dec.")])


# --------------------------------------------------------------------- the real workbook, damaged

def _damaged(parser, tmp_path, cell, value):
    raw, seeds = damage.copy_inputs(parser, tmp_path)
    damage.edit_cell(raw / "cbi" / KFI, "Sheet1", cell, value)
    return raw, seeds


def _august_2026_header():
    ws = openpyxl.load_workbook(REPO_ROOT / "pipeline" / "raw" / "cbi" / KFI, data_only=True, read_only=True)["Sheet1"]
    for col in range(3, ws.max_column + 1):
        if ws.cell(1, col).value == 2026:
            return f"{get_column_letter(col + 7)}2", ws.cell(2, col + 7).value  # January + 7 = August
    raise AssertionError("no 2026 block")


@pytest.mark.parametrize("parser", ["cbi_cpi", "cbi_money_banking"])
def test_misspelt_latest_month_stops_both_parsers(parser, tmp_path):
    cell, current = _august_2026_header()
    assert str(current).strip().lower().startswith("aug")
    raw, seeds = _damaged(parser, tmp_path, cell, "Augst")
    with block_network(), pytest.raises(ParseError) as info:
        damage.run_parser(parser, raw, seeds, tmp_path / "out")
    assert f"cell {cell}" in str(info.value) and "'Augst'" in str(info.value)
